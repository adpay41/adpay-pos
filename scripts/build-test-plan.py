"""
Builds docs/adpay-pos-test-plan.xlsx: the hand-testing workbook for a non-technical tester.

    python scripts/build-test-plan.py <register-url> <merchant-url> <admin-url>

The three URLs are the public links printed by `npm run share` (they change on every run).
Test cases marked checked=False were written from the code and not clicked through; the workbook
flags them so the tester knows an unexpected result there might be the plan's mistake.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "adpay-pos-test-plan.xlsx"

if len(sys.argv) != 4:
    sys.exit("usage: python scripts/build-test-plan.py <register-url> <merchant-url> <admin-url>")
REGISTER_URL, MERCHANT_URL, ADMIN_URL = sys.argv[1:4]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="111111")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=16, bold=True)
H2 = Font(name=FONT, size=12, bold=True)
MUTED = Font(name=FONT, size=10, italic=True, color="666666")
LINK = Font(name=FONT, size=11, color="0563C1", underline="single")
SECTION_FILL = PatternFill("solid", start_color="EDEDED")
FLAG_FILL = PatternFill("solid", start_color="FFF2CC")  # amber: expected result worked out from code
FILL_IN = PatternFill("solid", start_color="FFFFE0")  # pale yellow: the tester fills these in
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
FLAG_TEXT = "⚑ Worked out from the code, not yet clicked through by us. If it's different, report it anyway: the mistake may be ours.\n\n"
SCREEN_FILL = PatternFill("solid", start_color="F8CBAD")  # light red: a whole screen we have never clicked through
SCREEN_TEXT = "⚠ WE HAVE NOT CLICKED THROUGH THIS SCREEN AT ALL (only automated tests). Please try it carefully and report anything that looks wrong, confusing or broken, even small things.\n\n"


def steps(*lines: str) -> str:
    return "\n".join(f"{i}. {s}" for i, s in enumerate(lines, 1))


# Each case: (id, feature, steps, expected, checked). checked: True = we clicked it through; False = worked out from the code
# (amber); "screen" = a whole screen nobody has clicked through yet (light red, at the top of its sheet).
# A section header row is ("§", title); ("§!", title) is the red "do these first" header.
REGISTER = [
    ("§!", "⚠ DO THESE TWO FIRST: two register screens we have NOT clicked through ourselves. " + "If the register isn't signed in yet, do R1–R5 first (just below), then come back here."),
    ("T1", "Tax-free sale (whole ticket)", steps(
        "If the register isn't signed in yet, do R1–R5 first (just below), then come back here.",
        "Tap the Drinks tab, then Hot Coffee — Medium.",
        "In the bar along the bottom press Tax-free.",
        "Look at the box before choosing anything, then tap Resale certificate and type ST3-00417 in the certificate box. Press Make it tax-free.",
        "Approve: tap Luis Ortega, then 1 3 5 7.",
        "Press Tax-free again, press Charge tax again, approve as Luis again.",
        "Press Tax-free once more, choose Non-profit (no certificate), Make it tax-free, approve.",
        "Pay it by check so the drawer isn't started yet: press Check / other (under the quick-cash buttons), keep Check, press the black Take … button. Press Print receipt, read it, press Done."),
     "Before: Subtotal $2.25, Tax $0.15, Cash $2.40, Card $2.50. The Tax-free box offers Resale certificate, Non-profit, Government, Diplomat, Other exemption and a \"Certificate number (optional)\" box; \"Make it tax-free\" stays greyed until you pick a reason. It asks \"Manager approval: Make a sale tax-free\". After approval: the Tax-free button is dark, the tax row reads \"Tax (tax-free: Resale certificate)\" $0.00, Cash $2.25, Card $2.34. Charge tax again puts it back to Tax $0.15 / Cash $2.40. After the check: \"Sale complete\". The receipt says \"Tax exempt\" near the top of the totals, Tax $0.00, TOTAL $2.25, and \"Check $2.25\".", "screen"),
    ("T2", "Check and other tenders (EBT, gift card, house account)", steps(
        "If the register isn't signed in yet, do R1–R5 first (just below), then come back here.",
        "Tap the Sandwiches tab, then Chopped Cheese Hero. (Cash $11.72, Card $12.19.)",
        "On the keypad type 2 0 0 0 0 (that is $200.00) and press Check / other. Read the box, then press Back and press Clear on the keypad.",
        "Type 5 0 0 ($5.00), press Check / other. Tap EBT, type 4410 in the reference box, and press the black Take … button.",
        "With nothing typed, press Check / other again. Keep Check, type 1042 as the check number, press the black Take … button.",
        "Press Print receipt and read it, then Done."),
     "The box offers Check, EBT, Gift card, House account, Other tender and a reference box. With $200.00 typed it says \"Only $11.72 is left: a check or other tender can't be for more (no change)\" and the Take button is greyed. With $5.00: \"Pays $5.00 now; the rest on another tender.\" and the button reads \"Take $5.00 as EBT\"; afterwards the ticket shows \"Paid so far $5.00 · left $6.72 cash or $6.99 card\". The second time: \"Pays the rest: $6.72 (the cash price).\" and \"Take $6.72 as Check\", then \"Sale complete\". The receipt lists \"EBT #4410 $5.00\" and \"Check #1042 $6.72\" and \"Cash price applied\" (a check or EBT pays the cash price, never the card price).", "screen"),
    ("§", "Getting started (do these first, in order)"),
    ("R1", "Wrong setup code is refused", steps(
        "Open the Register link in Chrome. Make the window as large as you can (the register is built for a wide screen).",
        "In the code box type ZZZZ-0000 and press Pair register."),
     "Red text under the heading: \"That setup code is invalid, used or expired\". You stay on the \"Set up this register\" screen.", True),
    ("R2", "Pair the register to a store", steps(
        "In the code box, replace what you typed with JSQ3-DEMO.",
        "Press Pair register."),
     "The screen changes to \"Who's working?\" with \"Journal Square Deli & Grocery · Jersey City\" above it, a \"🌐 English\" button, "
     "and four people: Nadia Haddad (Owner), Luis Ortega (Manager), Dev Patel (Cashier), Maria Santos (Cashier).", True),
    ("R3", "Wrong PIN", steps(
        "Tap Maria Santos.",
        "Tap 1 1 1 1 on the keypad. There's no Enter key: it checks the PIN as soon as the 4th digit is in."),
     "\"Wrong PIN (4 tries left)\".", True),
    ("R4", "Five wrong PINs lock a person out", steps(
        "Tap ‹ Back, then tap Dev Patel.",
        "Tap 0 0 0 0 (no Enter). Wait a second for the message. Do this 5 times in total (wait each time; very fast taps get ignored).",
        "Tap ‹ Back."),
     "After the 5th try: \"Too many wrong PINs. Dev Patel is locked out on this register for 5 minutes.\" "
     "Back on \"Who's working?\", Dev Patel's tile says \"locked\".", True),
    ("R5", "Sign in: the new screen layout", steps("Tap Maria Santos, tap 2 4 6 8 (it signs in on the last digit)."),
     "The sale screen opens. Top bar: the store name, \"Register 3 (new)\", a line like \"Today $… · yesterday by now $…\", \"Clock out · 0:00\" (signing in starts Maria's shift), "
     "Maria / Lock, \"Drawer not started\", Synced, 🌐 English, Customer screen ↗. Under it, department TABS across the top: ★ Favorites, Sandwiches, Drinks, Snacks, "
     "Tobacco 21+, Lottery 18+, Grocery, Household. Then the search box and the item keys. Under the keys, the KEYPAD: a display, Clear, 7 8 9 ⌫ / 4 5 6 00 / 1 2 3 0, "
     "and @, PLU and Department keys; to its right Cash, Card, Check / other and Reject a bill. The ticket is on the right. Along the bottom: Void ticket, Reprint last, "
     "Hold, Discount, Return (no receipt), Tax-free, Tickets, Price check, Training, Receive, Write off, Checklist, End of day.", True),
    ("§", "Selling"),
    ("R6", "Ring an item", steps("Tap the Drinks tab, then tap Hot Coffee — Medium."),
     "The ticket shows Hot Coffee — Medium $2.25 with \"card $2.34\" under it. Subtotal $2.25, Tax $0.15. The two totals: Cash $2.40 and Card $2.50.", True),
    ("R7", "Tap twice = quantity 2", steps("Tap Hot Coffee — Medium again."),
     "Still one line, now \"2 × Hot Coffee — Medium\" $4.50 (card $4.68). Subtotal $4.50, Tax $0.30, Cash $4.80, Card $4.99.", True),
    ("R8", "Pay cash with one tap: start the drawer, quick cash, change, receipt", steps(
        "Look at the quick-cash buttons next to the keypad, then press $20.00. It asks you to \"Count the starting cash\" first (nobody started the drawer yet).",
        "Tap 1, 0, 0, 0, 0 (that is $100.00). Press \"Start drawer with $100.00\". Don't press anything else: the $20.00 goes through by itself.",
        "Press Print receipt, read it, then press Done."),
     "Quick buttons: Exact, $5.00, $10.00, $20.00, $50.00, and the Cash key reads \"Cash $4.80\" (there is no separate cash popup any more). After the drawer starts: "
     "\"Sale complete\", Change due $15.20, \"Total $4.80 cash · drawer opened\". "
     "The receipt (a box titled \"Receipt (printer preview)\") shows: 2 x Hot Coffee — Medium $4.50 · Subtotal $4.50 · Tax 6.625% on $4.50 $0.30 · TOTAL $4.80 · "
     "Cash price applied · Cash $20.00 · Change $15.20 · Cash price total $4.80 · Card price total $4.99, and near the bottom a web link with \"Scan for a digital copy\".", True),
    ("R9", "Search", steps("Click the search box (\"Search name, UPC or PLU — or scan\") and type cheet"),
     "The tile \"Flamin' Hot Cheetos 2.75 oz\" $2.29 appears. (Don't ring it.) Clear the search box.", True),
    ("R10", "Scan a barcode (typed)", steps("Click the search box, type 200000100131 and press Enter."),
     "Hot Coffee — Medium $2.25 is added to the ticket and the search box empties.", True),
    ("R11", "Unknown barcode becomes a new item", steps(
        "Click the search box, type 990011223344 and press Enter.",
        "In \"What is it?\" type Test Salsa 16oz. In \"Cash price\" type 3.49.",
        "Tap Grocery, then press Add & ring up."),
     "A \"New item\" box: \"Barcode 990011223344 isn't in the catalog yet…\". After typing the price, \"card $3.63\" appears next to it. "
     "After Add & ring up, the ticket has Test Salsa 16oz $3.49 (card $3.63).", True),
    ("R12", "Void an open ticket", steps("Press Void ticket."),
     "The ticket empties (\"Tap an item to start a sale.\"). No approval is asked.", True),
    ("R13", "Open-price item (do Admin case A6 first)", steps(
        "Tap Grocery, then Deli by weight (allow up to 15 seconds for it to appear after A6).",
        "Type 637 and press \"Ring up $6.37\".",
        "Press Void ticket afterwards."),
     FLAG_TEXT + "A keypad \"Price — Deli by weight\" appears. After ringing: Deli by weight $6.37, card $6.62.", False),
    ("R14", "Age-restricted item", steps(
        "Tap Tobacco, then Marlboro Red — Pack. Press Not verified.",
        "Tap it again and press \"ID checked — 21+\".",
        "Press Void ticket afterwards."),
     FLAG_TEXT + "First a \"Check ID — 21+\" box; Not verified adds nothing. After ID checked: Marlboro Red — Pack $14.00 with \"21+ checked\"; "
     "Cash $14.00 and Card $14.00 (tobacco has no card markup and no sales tax).", False),
    ("R15", "Hold and recall", steps(
        "Tap Sandwiches, then Buttered Roll. Press Hold.",
        "Press Held (1), then Recall.",
        "Press Void ticket."),
     FLAG_TEXT + "Hold empties the ticket and a black \"Held (1)\" button appears. The list \"Held tickets\" shows one ticket; Recall brings the Buttered Roll back.", False),
    ("R16", "Card payment (approved)", steps(
        "Ring Hot Coffee — Medium once (Drinks tab).",
        "Press the \"Card $2.50\" key next to the keypad. Wait a few seconds.",
        "Press No receipt."),
     FLAG_TEXT + "The card goes straight to the card machine (no extra Charge button): \"Card — $2.50\" and \"Customer: tap, insert or swipe on the card machine\". "
     "Then \"Sale complete\" and \"Paid $2.50 by card ···· 1234\" (the four digits change every time; that's normal).", False),
    ("R17", "Card declined (amounts ending in .13 always decline)", steps(
        "Tap the Grocery tab. Ring Breakfast Cereal and Whole Milk — Half Gallon.",
        "Press the \"Card $9.13\" key.",
        "Press \"Cash instead\", then tap Exact on the keypad, then No receipt."),
     FLAG_TEXT + "The ticket shows Cash $8.78 and Card $9.13 (both items say \"no tax\"). After Card: \"Declined — test card: amounts ending in .13 decline. "
     "Try again, another card, or cash.\" Cash instead closes the card box; Exact completes the sale with change $0.00.", False),
    ("R18", "Split: part cash, part card", steps(
        "Ring Hot Coffee — Medium twice.",
        "On the keypad tap 2, 00 (that is $2.00). The Cash key now reads \"Part cash $2.00\": press it.",
        "Press \"Charge $2.91\", then No receipt."),
     FLAG_TEXT + "The card box: \"Paid so far $2.00\", \"Left to pay by card $2.91\", \"or $2.80 in cash\". Then \"Sale complete\" and \"Paid $4.91: $2.00 cash + $2.91 card\".", False),
    ("R19", "Price check", steps(
        "Press Price check.",
        "Tap Sandwiches → Chopped Cheese Hero.",
        "Press \"Show cost and margin\", then Cancel.",
        "Press Done, then Price check again to turn it off."),
     FLAG_TEXT + "A banner \"Price check — scan or tap an item to see its prices. Nothing is rung up.\" A box \"Price check — not rung up\" with Cash $10.99 and Card $11.43. "
     "Show cost asks for a manager/owner PIN (Maria can't see costs). The ticket stays empty.", False),
    ("R20", "Reprint the last receipt", steps("Press Reprint last."),
     FLAG_TEXT + "A receipt preview of the last sale, marked *** REPRINT ***.", False),
    ("R21", "Refund needs a manager", steps(
        "Press Tickets. Tap the $2.50 card sale from R16.",
        "Press + next to Hot Coffee — Medium (reason \"Returned\" is already selected).",
        "Press \"Refund $2.50 to card\".",
        "Tap Luis Ortega and tap 1 3 5 7."),
     "\"Who's approving?\" lists Nadia Haddad and Luis Ortega. After the PIN, a receipt preview appears with \"Refunded -$2.50\" and at the bottom "
     "\"REFUND - $2.50 returned\".", True),
    ("R22", "Void a completed sale", steps(
        "Press Tickets and open the $4.80 cash sale from R8.",
        "Press Void sale. Approve as Nadia Haddad (2 5 8 0)."),
     FLAG_TEXT + "The sale shows as voided and \"Drawer opened\" flashes (the $4.80 goes back in cash).", False),
    ("R23", "Training mode", steps(
        "Press Training.",
        "Ring Hot Coffee — Medium, tap Exact on the keypad → Print receipt.",
        "Press Exit training, then Tickets."),
     FLAG_TEXT + "A black banner \"TRAINING — practice only…\". Card, Tickets and End of day disappear. The receipt ends \"*** TRAINING — NOT A SALE ***\". "
     "After leaving training, the practice sale is not in Tickets.", False),
    ("R24", "Drawer: safe drop, paid-out, no sale", steps(
        "Tap Drawer (top bar).",
        "Safe drop: tap 5,0,0,0 ($50.00), Next, choose Safe drop, press \"Open drawer & record\".",
        "Paid out: approve as Luis (1357), tap 2,0,0,0 ($20.00), Next, choose Vendor delivery, payee Stella Bakery, \"Open drawer & record\".",
        "Press No sale and approve as Luis."),
     FLAG_TEXT + "The Drawer panel says it was started by Maria Santos with $100.00. After the drop: Drops $50.00. Paid out asks for approval, then Paid out $20.00. "
     "No sale asks for approval, then \"Drawer opened\".", False),
    ("R25", "Close the drawer (blind count)", steps(
        "In Drawer, press \"Close & count drawer\".",
        "Tap 2, 5, 0, 0 ($25.00) and press \"Close with $25.00 counted\".",
        "Write down the Expected and Short amounts (you need them in X3)."),
     "Before you count, Maria is NOT shown what the drawer should hold. After closing: \"Drawer closed\", Counted $25.00, Expected $X, and \"Short $Y\" in amber (never red), "
     "where Y = X − 25.00. X is $100.00 plus every cash payment you took since R8, minus drops and paid-outs (it depends on what you did). "
     "\"Recorded for Maria Santos. The owner sees this in the merchant app.\"", True),
    ("R26", "End of day (Z-report)", steps("Press End of day, then \"Take Z & print\"."),
     "First: \"Since the last Z: N sales, $… (cash …, card …), tax …\". Then a printout \"Z-REPORT #…\" (#1 on the fresh demo, higher if someone already took one) with "
     "Sales, BY CATEGORY, TAX, Refunds, Voids, No-sale opens, Counterfeits refused, and CASH DRAWER (Float, Expected, Counted $25.00, Short from R25).", True),
    ("R27", "Selling with no internet", steps(
        "Press F12 (Chrome DevTools) → Network tab → change \"No throttling\" to Offline. Don't reload the page.",
        "Ring Hot Coffee — Medium. Tap Exact on the keypad (start the drawer with 10000 if asked), then No receipt.",
        "Set DevTools back to No throttling and wait 15 seconds."),
     FLAG_TEXT + "A yellow banner \"Cash only right now — no connection…\", and Card shows \"offline\". The sale completes; the status pill says \"Offline · N queued\". "
     "After reconnecting the pill goes back to Synced.", False),
    ("R28", "Lock and switch cashier", steps("Press Lock (top right). Sign in as Luis Ortega (1 3 5 7)."),
     FLAG_TEXT + "Back to \"Who's working?\", then the sale screen with \"Luis\" at the top.", False),
    ("§", "Customer screen, languages and loyalty"),
    ("R29", "Customer screen", steps(
        "Press \"Customer screen ↗\". Allow pop-ups if Chrome asks. Put the new window next to the register.",
        "Look at it before ringing anything."),
     "The customer window shows the store name, a 🌐 English button, \"Larger text\", \"Cash and card prices are both shown before you pay.\" and the deal line "
     "\"★ 2 for $6.00 · Energy drinks\".", True),
    ("R30", "Customer picks Spanish", steps("On the customer window, press 🌐 English and choose Español."),
     "Only English and Español are offered. The customer window switches to Spanish: \"Letra más grande\", \"Los precios en efectivo y con tarjeta se muestran antes de pagar.\" "
     "(The deal line stays \"2 for $6.00 · Energy drinks\"; see Don't Report This.)", True),
    ("R31", "Loyalty by phone (customer side)", steps(
        "On the register ring Hot Coffee — Large (Drinks).",
        "On the customer window tap \"★ Gane recompensas — ingrese su número de teléfono\".",
        "Tap 2 0 1 5 5 5 0 1 2 3 and press Listo. Don't tick the texts box."),
     "Before the phone: both prices, \"Pagar en efectivo $2.93\" and \"Pagar con tarjeta $3.05\". After Listo the customer window says "
     "\"Teléfono terminado en 0123 · 0 de 5 visitas — 5 más para su recompensa\" (if someone already used 0123, the visit count is higher). "
     "The register ticket header says \"CUSTOMER: SPANISH\" and \"Rewards ···0123 · 0/5 visits\".", True),
    ("R32", "Receipt in the customer's language", steps("On the register tap Exact on the keypad, then Print receipt. Press Done."),
     "\"Sale complete … receipt in Spanish\". The customer window says \"¡Gracias!\" and \"Pagado $2.93\". The receipt shows Impuesto 6.625% sobre $2.75 $0.18, "
     "Precio en efectivo aplicado, Efectivo $2.93, Cambio $0.00, Escanee para una copia digital. The store's own lines (returns policy, thank you) stay in English.", True),
    ("R33", "Customer 'Larger text'", steps("On the customer window press \"Letra más grande\" (Larger text)."),
     FLAG_TEXT + "The customer window's text gets noticeably bigger. Press it again to go back.", False),
    ("R34", "Cashier's own language", steps(
        "On the register press 🌐 English (top bar) and choose Español.",
        "Look at the buttons, then switch back to English the same way."),
     "A menu lists English, Español, 中文, 한국어, العربية, हिन्दी. In Spanish the register shows Efectivo, Tarjeta, Cierre del día, Lista de tareas, Recibir, Dar de baja. "
     "Item and category names stay as the store typed them.", True),
    ("R35", "Digital receipt link", steps("Copy the web link printed on any receipt (above \"Scan for a digital copy\") into a new Chrome tab."),
     "A page shows the same receipt as the printout.", True),
    ("§", "Deals, stock and store routines"),
    ("R36", "Deal: any 2 energy drinks for $6.00", steps(
        "Tap Drinks. Ring Red Bull 8.4 oz, then Monster Energy 16 oz.",
        "Tap Exact on the keypad → Print receipt."),
     "Each line gets \"Energy drinks −$…\" under it (−$0.57 and −$0.61). Subtotal $6.00, Tax $0.40, Cash $6.40, Card $6.65. "
     "The receipt shows the two \"Energy drinks\" lines and \"You saved $1.18\".", True),
    ("R37", "Low stock and sell-soon", steps("Look at the Monster Energy tile, and at the ticket area when it's empty."),
     "The Monster tile has a small \"N left\" badge (N is 12 or less; it changes as people sell). An empty ticket shows an amber line "
     "\"Sell soon: Whole Milk — Gallon (…) by …\" (the quantity and date change).", True),
    ("R38", "Receive a delivery", steps(
        "Press Receive.",
        "Invoice number: BG-1001. In \"No barcode? Type the name\" type red bull and tap Red Bull 8.4 oz in the list.",
        "Press + until the quantity is 12. Press \"Add 12 to stock\"."),
     "A \"Receive a delivery\" box. After adding: a box titled \"Done\" with \"Received 12 items.\" Press OK.", True),
    ("R39", "Write off spoiled stock (needs a manager)", steps(
        "Press Write off. Approve as Luis (1 3 5 7).",
        "In \"Which item?\" type whole milk and tap Whole Milk — Gallon. Set How many to 2. Tap Spoiled. Press Write off."),
     "\"Who's approving?\" appears first. Then a box titled \"Done\": \"Wrote off 2 × Whole Milk — Gallon.\"", True),
    ("R40", "Closing (or opening) checklist", steps(
        "Press Checklist. (Before noon it opens on the Opening checklist, after noon on Closing.)",
        "On the Closing checklist tick \"Floors swept, trash out\" and \"Back door locked\".",
        "In \"Note for the owner\" type Slicer blade needs replacing. Press Save checklist."),
     "The closing list has: Deli slicer cleaned (Photo), Floors swept, trash out, Cash dropped in the safe (Photo), Back door locked, Lights and signs off, alarm on. "
     "After saving: a box titled \"Done\" with \"Closing checklist saved: 2 of 5 done.\"", True),
    ("R41", "Repeat last sale", steps("With an empty ticket, press ↻ Repeat last."),
     FLAG_TEXT + "The items of the previous sale are rung again. Void the ticket afterwards.", False),
    ("§", "New: the keypad, departments, keys, discounts and returns"),
    ("R42", "Quick cash in one tap", steps("Tap the Sandwiches tab, then Buttered Roll.", "Press $5.00 next to the keypad. Press No receipt."),
     "The quick buttons read Exact, $3.00, $5.00, $10.00, $20.00; the Cash key \"Cash $2.12\" and the Card key \"Card $2.21\". One tap on $5.00: \"Sale complete\", Change due $2.88.", True),
    ("R43", "Pay with a typed amount", steps("Tap the Drinks tab, then Hot Coffee — Large.", "On the keypad tap 1, 00, 0.", "Press the Cash key, then No receipt."),
     "The display shows $10.00 and the Cash key changes to \"Cash $10.00\". The sale ($2.93) completes with Change due $7.07.", True),
    ("R44", "Ring an amount to a department (no item)", steps("Tap the Grocery tab.", "On the keypad tap 3, 5, 0.", "Press the \"→ Grocery\" key (next to the digits). Then Void ticket."),
     "The department key reads \"→ Grocery\" and under it $3.50. The ticket gets a line \"Grocery\" $3.50 \"no tax\" (Grocery isn't taxed); the keypad clears. "
     "On ★ Favorites the key says \"Department · pick a tab\" and can't be pressed.", True),
    ("R45", "Department amount with an age check", steps("Tap the Tobacco tab. Tap 1, 2, 00 on the keypad. Press \"→ Tobacco\".", "Press \"ID checked — 21+\". Then Void ticket."),
     FLAG_TEXT + "A \"Check ID — 21+\" box first. After ID checked: a line \"Tobacco\" $12.00 with \"21+ checked\", no tax, and you are not asked for the price again.", False),
    ("R46", "@ key: a number times the next item", steps("Tap the Drinks tab. On the keypad tap 3, then press @.", "Tap Hot Coffee — Medium.", "Tap Hot Coffee — Medium once more. Then Void ticket."),
     "After @ the display shows \"3 ×\" and \"Now tap, scan or ring the item\". The coffee comes in as \"3 × Hot Coffee — Medium\" $6.75. "
     "The next tap adds just one (the line becomes 4 ×): the 3 × is used once.", True),
    ("R47", "PLU key (do Admin case A23 first)", steps("Tap 4, 0, 1, 1 on the keypad and press PLU.", "Tap 0, 4, 0, 1, 1 and press PLU.", "Tap 7, 7, 7 and press PLU. Then Void ticket."),
     FLAG_TEXT + "Banana $0.39 is rung; the second time (with a leading zero) it rings Banana again (the line becomes 2 ×). 777 shows \"No item with PLU 777.\"", False),
    ("R48", "Your own key page (do Merchant case MP1 first)", steps("Look at the tabs, then tap Coffee bar.", "Tap the \"$10 grocery\" key, then Void ticket."),
     "A \"Coffee bar\" tab sits right after ★ Favorites, before the departments. It shows your keys in the order you left them: the three coffees as normal keys and "
     "\"$10 grocery\" with GROCERY, $10.00 and card $10.40. The $10 key rings a line \"Grocery\" $10.00.", True),
    ("R49", "Discount the whole ticket (needs a manager)", steps(
        "Ring Buttered Roll (Sandwiches) and Hot Coffee — Large (Drinks).",
        "Press Discount (bottom bar). Tap 10%, tap Regular customer, press \"Take $0.47 off\". Approve as Luis (1 3 5 7).",
        "Tap Exact on the keypad, then Print receipt. Press Done."),
     "Before: Subtotal $4.74, Tax $0.31, Cash $5.05. The box says \"Takes $0.47 off the cash price (card price in proportion).\" After approval the lines still show $1.99 and $2.75, "
     "a row \"Ticket discount −$0.47\", Subtotal $4.27, Tax $0.28, Cash $4.55, Card $4.73; the Discount button reads \"Discount on\". The receipt shows \"Discount 10% -$0.47\", Subtotal $4.27, TOTAL $4.55.", True),
    ("R50", "A price that already includes tax (do Merchant case M33 first)", steps(
        "Tap the Drinks tab, then Hot Coffee — Large.",
        "Tap the Sandwiches tab, then Buttered Roll.",
        "Tap Exact on the keypad, then Print receipt. Press Done. Then undo M33 in the merchant app."),
     FLAG_TEXT + "With just the coffee: Subtotal $2.75, Tax $0.00, \"Tax included in prices $0.17\", Cash $2.75, Card $2.86 (no tax added on top). With the roll too: Subtotal $4.74, "
     "Tax $0.13, Tax included in prices $0.17, Cash $4.87. The receipt: \"Hot Coffee — Large *\" $2.75, Subtotal $4.74, \"Tax 6.625% on $1.99\" $0.13, TOTAL $4.87, "
     "\"* Incl. tax 6.625% on $2.58\" $0.17, \"Total tax\" $0.30, and \"* Price includes tax\".", False),
    ("R51", "Refund without a receipt (needs a manager)", steps(
        "With an empty ticket press Return (no receipt) (bottom bar). Tap Defective, press Start the return. Approve as Luis (1 3 5 7).",
        "Tap the Sandwiches tab, then Buttered Roll.",
        "Press the black \"Refund $2.12 in cash\" button."),
     "It asks \"Manager approval: Refund without a receipt\". Then the ticket shows a black banner \"RETURN — no receipt · Defective\" and the keypad's Cash/Card keys are off. "
     "After the roll: \"Refund $2.12 in cash\" ($1.99 + tax, at today's price). After pressing it: a box titled Done, \"Refunded $2.12 in cash.\", and a new empty ticket.", True),
]

MERCHANT = [
    ("§!", "⚠ DO THESE FIRST: a merchant-app screen we have NOT clicked through ourselves (Items → Pages). Sign in as in M1 first (just below)."),
    ("MP1", "Build a key page for the register", steps(
        "Sign in as Nadia (see M1). Open Items, then the Pages section.",
        "Press + Page. In \"Page name\" replace the text with Coffee bar.",
        "Under \"Add an item\" type coffee and tap + Hot Coffee — Small, + Hot Coffee — Medium, + Hot Coffee — Large.",
        "Under \"Add a department amount\" tap Grocery, type 10.00 as the amount and $10 grocery as the label. Press Add.",
        "On the \"$10 grocery\" key press ↑ once. Press Save key pages.",
        "Look at the register within about 15 seconds (then do Register case R48)."),
     "A \"Key pages\" card with a chip \"Coffee bar (4)\" after you add the keys. The keys list reads \"Hot Coffee — Small · $1.75\", \"… Medium · $2.25\", "
     "\"$10 grocery · Grocery · $10.00\" (after ↑ it sits above Large), \"… Large · $2.75\". After saving: \"Key pages saved (1). Registers update within 15 seconds.\" "
     "The register gets a Coffee bar tab after ★ Favorites with the keys in that order.", "screen"),
    ("MP2", "Key pages: limits and mistakes", steps(
        "Still on Pages: press + Page, name it coffee BAR (same name, different capitals). Press Save key pages.",
        "Delete that page (Delete page) and save again. Try the ← Earlier / Later → buttons and × on a key; don't save those changes (sign out and back in to undo)."),
     "Saving two pages with the same name is refused with \"Two pages have the same name\" and nothing changes on the register. After deleting the duplicate, the save works. "
     "The Save button is greyed whenever nothing has changed.", "screen"),
    ("§", "Signing in and the owner's numbers"),
    ("M1", "Sign in with a code", steps(
        "Open the Merchant app link.",
        "Type 2015550100 and press \"Text me a code\".",
        "Type 123456 and press Enter."),
     "\"ENTER THE 6-DIGIT CODE\" with \"Local dev — no SMS is sent. Your code: 123456\". Then the store \"Journal Square Deli & Grocery\", \"Nadia Haddad\", and tabs: "
     "Sales, Tickets, Cash, Hours, Customers, Items, Stock, Orders, Deals, Alerts, Staff, Help.", True),
    ("M2", "Sales today, live ticker, roll-up", steps("On Sales → Today, scroll down slowly."),
     "Cards in this order: ALL YOUR STORES (Jersey City and Astoria, each with tickets and $), PROFIT, TODAY SO FAR (vs yesterday and last week), LIVE with \"● connected\" "
     "and recent sales with register and cashier, NET SALES, CARD, CASH, BY HOUR, BY REGISTER, BY CASHIER, TAX & COMPLIANCE with three Export buttons. "
     "Amounts change as people sell.", True),
    ("M3", "7 days and Month", steps("Switch Today → 7 days → Month."),
     FLAG_TEXT + "The numbers change; nothing errors.", False),
    ("M4", "Tickets", steps("Open Tickets."),
     FLAG_TEXT + "Recent tickets with time, \"Jersey City\", register, item count and total. Your register sales are there; the R22 void shows as voided.", False),
    ("M5", "Cash: over/short", steps("Open Cash → Today (after register case R25)."),
     "OVER / SHORT shows the shortage from R25 (in amber, not red). BY CASHIER: Maria Santos with the same short. DRAWER SESSIONS: the Register 3 session with "
     "\"Maria Santos · float $100.00\" and the short. END OF DAY lists your Z from R26. Your closed drawer is NOT listed under \"In the drawers now\".", True),
    ("M6", "Hours and payroll export", steps(
        "On the register press Lock, then sign in as Dev Patel (3 6 9 0): signing in starts his shift.",
        "In the app open Hours → This week, then press \"Export for payroll (CSV)\"."),
     FLAG_TEXT + "Dev Patel shows \"on the clock now\". A CSV file downloads.", False),
    ("M7", "Opening/closing checklist report", steps("Open Hours and scroll to \"Open & close checklists\"."),
     "A row for today, Jersey City: Opening \"not done\" (unless someone did it), Closing \"2/5 · 2 photos missing\" with the time and \"Maria Santos\". "
     "Tap the Closing entry: the five items with ✓/✗, \"no photo\" on the two photo items, and your note in quotes.", True),
    ("M8", "Edit the checklist", steps(
        "In \"Open & close checklists\" press Edit lists.",
        "Under Opening checklist press + Add item and type Check the ATM has paper. Press Save lists."),
     FLAG_TEXT + "The editor closes. On the register, press Checklist → Opening checklist (within about 15 seconds): the new item is at the bottom.", False),
    ("M9", "Cashier performance", steps("On Hours, scroll to CASHIER PERFORMANCE."),
     "One block per person: sales $, number of sales, average, time on the clock, $ per hour, voids, refunds, \"no sale\" opens, drawer over/short, age checks. "
     "The numbers depend on what everyone did.", True),
    ("§", "Customers and loyalty"),
    ("M10", "Loyalty program and customer list", steps("Open Customers."),
     "\"Loyalty program\" with Punch card (visits) chosen, reward every 5 visits, WHAT COUNTS: Drinks, REWARD \"A free item (Drinks)\". Below: \"CUSTOMERS (…)\" "
     "including \"Phone ···0123\" from register case R31 with its visit count and total. Only the last four digits are ever shown.", True),
    ("M11", "Text a deal to regulars", steps("In Customers, find \"Text a deal to your regulars\"."),
     FLAG_TEXT + "It says how many customers said yes to texts (0 unless someone ticked the box). Sending is recorded but no text is delivered (see Don't Report This).", False),
    ("§", "Items and prices"),
    ("M12", "Price history", steps("Open Items → Items and tap Snickers."),
     "\"Edit item\" with the fields and, at the bottom, PRICE HISTORY with at least one line like \"Sep 27, 2026 · $1.99 · import\" (the date the demo data was loaded).", True),
    ("M13", "Change a price, see it in history and on the register", steps(
        "On Snickers change CASH PRICE to 2.19 and press Save.",
        "Open Snickers again, then check the register (Snacks).",
        "Put it back to 1.99 and Save."),
     FLAG_TEXT + "A new PRICE HISTORY line $2.19 at the top. Within about 15 seconds the register's Snickers tile shows $2.19 (card $2.28), without reloading.", False),
    ("M14", "Bulk price change: preview only", steps(
        "Items → Bulk. Tap Snacks. Keep \"± %\", replace the 5 with 10. Tap \"Round up to .99\". Press Preview.",
        "Do NOT press \"Change 15 prices\"."),
     "\"15 prices change.\" and a list including Beef Stick: $1.49 → $1.99, Butterscotch Krimpets: $1.99 → $2.99, Pringles Original: $2.99 → $3.99. "
     "(If someone changed snack prices, the list differs.)", True),
    ("M15", "Shelf tags (PDF)", steps("Items → Tags. Tap \"Snacks (15)\" under OR EVERY ITEM IN A CATEGORY. Allow pop-ups if asked."),
     "The page lists items whose price changed since their tag was printed and a \"Print N tags\" button (N varies). "
     "Tapping Snacks opens a PDF of 15 tags in a new tab, each with the cash and card price and a barcode.", True),
    ("M16", "Add an item with a photo", steps(
        "Items → Items → + Add item. Name Test Muffin, category Sandwiches, cash price 2.49.",
        "Press Choose photo and pick any picture from your computer. Press Save."),
     FLAG_TEXT + "The card price line says it's automatic (cash + 4% = $2.59). After saving, within about 15 seconds the register's Sandwiches page shows a Test Muffin tile with the photo.", False),
    ("M17", "Arrange favorite keys", steps("Items → Favorites → Arrange keys. Tap the first key, then the last. Press Save favorites."),
     FLAG_TEXT + "The first key moves to the end; the register's ★ Favorites page follows within about 15 seconds.", False),
    ("§", "Deals"),
    ("M18", "Create a deal and try it", steps(
        "Open Deals → New deal. Keep \"N for $X\".",
        "In the first box delete what's there and type 2; in the \"for $\" box delete what's there and type 2.00.",
        "Name: Beef sticks 2 for $2. In \"…or add items\" type beef and tap Beef Stick.",
        "Press Save deal. Wait 15 seconds.",
        "On the register scan 200000100421 twice (type it in the search box, Enter, twice)."),
     "Under the form a green preview \"★ 2 for $2.00 · Beef sticks 2 for $2\", then \"Saved. Registers pick it up at their next sync.\" On the register: "
     "\"2 × Beef Stick\" with \"Beef sticks 2 for $2 −$0.98\", Subtotal $2.00, Tax $0.13, Cash $2.13, Card $2.22. Void the ticket. "
     "", True),
    ("M19", "End a deal", steps("In Deals, press \"End now\" under your Beef sticks deal."),
     "It changes to \"Ended\" with a \"Run again\" button. Its line \"Last 30 days: 0 tickets\" (voided tickets don't count). The Energy drinks deal keeps running.", True),
    ("§", "Stock and ordering"),
    ("M20", "Stock levels and shrink", steps("Open Stock."),
     "SELL SOON with the milk lot, then All (5) / Low / Not selling filters and five tracked items with their counts (Monster is marked \"Low: at or below 12\"), "
     "then SHRINK · LAST 30 DAYS including your write-off (\"Write-offs: spoilage …\") and a BY CASHIER list. Counts change as people sell and receive.", True),
    ("M21", "Order from a vendor", steps(
        "Open Orders. Under \"Order from Tuscan Dairy\" press \"Make the order\".",
        "Under ORDERS, press \"Send to the rep\"."),
     "\"Order for Tuscan Dairy saved as a draft…\", the draft lists \"Whole Milk — Gallon: ordered N\" (N depends on sales). After sending: "
     "\"Recorded for •••0173, but not delivered: texting and email aren't connected yet. Call or text the rep the order.\" and the order says \"sent (not delivered)\".", True),
    ("§", "Alerts, staff, help"),
    ("M22", "Alerts", steps("Open Alerts."),
     "A list that includes \"Tobacco retail licence (demo) expires in N days\", your drawer short from R25, and a Dev Patel lockout from R4 (plus older alerts). "
     "REGISTERS shows Register 3 (new) · Jersey City · online.", True),
    ("M23", "Alert settings", steps("Open ALERT SETTINGS ▸, switch off \"Large refund or void\", press Save alert settings. Switch it back on and save."),
     FLAG_TEXT + "\"Saved.\" both times.", False),
    ("M24", "Add a staff member", steps("Staff → + Add a person. Name Test Cashier, role Cashier, PIN 9876, Add. Then on the register press Lock."),
     FLAG_TEXT + "Test Cashier is listed; within about 15 seconds the register's \"Who's working?\" shows Test Cashier, who can sign in with 9876.", False),
    ("M25", "Report a problem to AD Pay", steps(
        "Help → Problems & equipment. What's wrong: Card machine beeps twice. Details: Register 3, since this morning.",
        "Press Send to AD Pay."),
     "\"Sent. AD Pay answers within one business day; urgent problems, call us.\" Your ticket appears as \"With AD Pay\" under the demo ticket \"Receipts come out faint\".", True),
    ("M26", "Your equipment", steps("Stay on Help → Problems & equipment and scroll down."),
     "\"Your equipment\" lists 9 units: a printer, a register and a card terminal for each of Registers 1, 2 and 3, each with a DEMO serial and \"warranty to …\".", True),
    ("M27", "Documents", steps("Help → Documents."),
     "Two documents: \"Tobacco retail licence (demo)\" with \"Expires in N days\" in amber, and \"General liability insurance (demo)\" \"Expires 2027-04-14\"; "
     "each with View, Upload renewal and Remove.", True),
    ("M28", "Upload a renewal", steps(
        "Under the tobacco licence press Upload renewal. Expires on: 2027-10-31.",
        "Choose any PDF or photo from your computer. Press Save."),
     FLAG_TEXT + "\"Renewal saved; the old one is kept in the history.\" The licence now says \"Expires 2027-10-31\" and \"1 earlier version\". "
     "Within a couple of minutes the \"expires in N days\" alert disappears from Alerts.", False),
    ("M29", "Support chat", steps("Help → Chat. Type Hello from the tester and press Send."),
     FLAG_TEXT + "Your message appears in the conversation (AD Pay sees it in Admin → Support).", False),
    ("§", "A different business (Bayonne)"),
    ("M30", "Another business can't see this one", steps(
        "Press Sign out. Sign in with 2015550142 and code 123456.",
        "Open Tickets and Items."),
     "Store \"Bayonne Corner Mart\", person Kevin Walsh, no way to switch to Journal Square. Tickets show only Broadway sales. Items have no Test Muffin or Test Salsa.", True),
    ("M31", "Product library suggests a name (see Cross-App X8 for the full flow)", steps(
        "Still as Kevin: Items → + Add item. Name Test Tissues 2-Ply, any category, cash price 2.49, BARCODE (UPC) 036000291452. Save.",
        "Sign out and sign in as 2015550100 (Nadia). Items → + Add item. In BARCODE (UPC) type 036000291452 (don't save)."),
     FLAG_TEXT + "Under the barcode: \"1 store on AD Pay calls it “Test Tissues 2-Ply”.\" and the Name box fills in with Test Tissues 2-Ply (Kevin's store has it). "
     "Leave without saving.", False),
    ("§", "Tax by category and prices that include tax"),
    ("M32", "Switch a category between Taxed and No tax", steps(
        "Items → Categories. Look at the Taxed / No tax chip on each category.",
        "Tap Snacks' Taxed chip so it says No tax. Within 15 seconds ring Snickers on the register (Snacks tab).",
        "Void the register ticket and tap Snacks' chip back to Taxed."),
     FLAG_TEXT + "Grocery, Tobacco and Lottery say No tax; the others Taxed. With Snacks on No tax the register rings Snickers $1.99 with \"no tax\" and Tax $0.00. "
     "After switching back, a new Snickers rings with tax again.", False),
    ("M33", "Mark an item as \"price includes tax\" (used by Register R50)", steps(
        "Items → Items → tap Hot Coffee — Large.",
        "Switch on \"Price includes tax\" and read the line under it. Press Save.",
        "Do Register case R50. Afterwards switch it off again and Save."),
     FLAG_TEXT + "Switched on, the line under it reads \"Rings at exactly this price. The tax inside it is still counted on your tax report.\" (off: \"Tax is added on top of this price…\"). "
     "After saving, the item list shows Hot Coffee — Large with \"· tax included\".", False),
    ("§", "Bringing a store over from NRS (do these after the Register sheet: they add new tabs)"),
    ("M34", "Save the sample NRS file", steps("Follow the steps on the sheet \"NRS sample file\" to save nrs-sample.json on your computer."),
     "You have a file nrs-sample.json of about 5 KB. (A real store's file comes from its NRS login; this is a made-up 12-item one.)", True),
    ("M35", "Preview an NRS price book", steps(
        "Items → Import. Press \"Choose the NRS file\" and pick nrs-sample.json.",
        "Press Preview and read everything on the page. Don't import yet."),
     FLAG_TEXT + "\"Preview — nothing saved yet\" with: \"Will import 12 of 12 items: 12 new, 0 updated, 0 unchanged.\", "
     "\"Categories: 5 NRS departments, 4 new, 1 already in your catalog (their tax and age settings kept).\", \"Barcodes: 11 real product barcodes, 1 store codes, 0 without a barcode.\", "
     "\"Quick keys: 0 — …\", \"Tax-inclusive prices: 1 items ring at their marked price, tax inside.\" Under COULDN'T CARRY OVER EXACTLY: 1 × open price, 1 × store barcode, "
     "1 × short code, 11 × no cost. Under DEPARTMENTS → CATEGORIES: CHECK TAX AND AGE: Drinks (already in your catalog, settings kept), Grocery tax (Taxed), Grocery Non-Taxable (No tax), Beers (Taxed, Alcohol 21+), "
     "all smoke (Taxed, Tobacco 21+), each with Taxed / No tax and age chips you can change. ITEMS TO LOOK AT AFTER IMPORTING repeats those with examples, then FIRST CHANGES with \"New · …\" lines and an \"Import 12 items\" button.", False),
    ("M36", "Import it", steps("Leave the department chips as they are and press Import 12 items. Wait for it to finish."),
     FLAG_TEXT + "\"Imported\" and \"Imported 12 of 12 items: 12 new, 0 updated, 0 unchanged.\", plus a notice that registers update within 15 seconds. "
     "Items now lists e.g. Coca-Cola Classic 20oz $2.49 and Marlboro Red Box $13.99 · tax included, and Gatorade Strawberry Kiwi 28oz as open price. Then do Cross-App X14.", False),
    ("M37", "Uploading again never makes duplicates", steps("Choose the same file again and press Preview. Don't import."),
     FLAG_TEXT + "\"Will import 12 of 12 items: 0 new, 0 updated, 12 unchanged.\" and Categories \"… 0 new, 5 already in your catalog …\".", False),
]

ADMIN = [
    ("§", "Signing in"),
    ("A1", "Sign in", steps("Open the Admin link. Email admin@adpay.local, password adpay-demo. Press Sign in."),
     "The Merchants page with Bayonne Corner Mart and Journal Square Deli & Grocery (Astoria and Jersey City). Menu: Merchants, Onboarding, Fleet, Alerts, Support, "
     "Tickets, Hardware, Rollouts, Partners, Agents, Investor pack, UPC library, Sales, Money, Translations, Audit log. (There is no Tax page any more: see A5.)", True),
    ("A2", "Wrong password", steps("Sign out and try the password wrong."),
     FLAG_TEXT + "An error; you're not signed in. Sign back in correctly.", False),
    ("§", "Stores, catalog and setup"),
    ("A3", "Merchant page tabs", steps("Open Journal Square Deli & Grocery. Click each tab: Sales, Cash, Catalog, Staff, Plan & setup."),
     FLAG_TEXT + "Each tab opens without an error; Sales → Latest tickets includes your register sales. Plan & setup includes \"Agent / referral\" and \"Documents\" panels.", False),
    ("A4", "Edit a price (used in X1)", steps("Catalog tab. Open Snickers. Change cash to 2.09 and press Save changes."),
     FLAG_TEXT + "Cash $1.99 and card \"Automatic: cash + 4%\" before; a green saved message with a new catalog version after.", False),
    ("A5", "Tax schedule", steps(
        "Catalog tab, set \"Priced at location\" to Astoria. Tax & compliance at Astoria → Edit → + Add rate. Class standard, rate 9, date 2099-01-01. Save."),
     FLAG_TEXT + "The Tax & compliance panel for Astoria shows sales tax today 8.875% and, scheduled, \"standard 9% from 2099-01-01\". (Whether a category is taxed is set by the "
     "store owner in the merchant app: Merchant case M32.)", False),
    ("A6", "Open-price item (needed for R13)", steps("Catalog (Jersey City) → + New item. Name Deli by weight, category Grocery, tick Open price, leave cash empty. Add item."),
     FLAG_TEXT + "The item list shows Deli by weight with price \"open\".", False),
    ("A7", "Onboard a new store, with a referral code (do A18 first)", steps(
        "Onboarding → + New merchant. Store name Tester Deli; owner Test Owner, phone 9175550188; state NJ; pricing defaults; install date tomorrow; "
        "Referral code TEST01.",
        "On Review press Create merchant, then \"print the install kit\" → Issue codes."),
     FLAG_TEXT + "\"Merchant created with 1 register\"; Tester Deli in the pipeline: Setting up, KYB not started, 0 / 1 paired. The install kit shows a QR code and a setup code. "
     "Agents → your agent now lists Tester Deli.", False),
    ("A8", "Documents on the merchant page", steps("Merchants → Journal Square → Plan & setup → Documents → Open next to the tobacco licence. Allow pop-ups."),
     FLAG_TEXT + "The two documents with their expiry pills; Open shows the file in a new tab (and is recorded in the Audit log).", False),
    ("A9", "Translations", steps("Open Translations."),
     FLAG_TEXT + "A list of languages with their status (Spanish offered; the others draft) and coverage.", False),
    ("§", "Support and fleet"),
    ("A10", "Tickets with SLA and a canned fix", steps(
        "Open Tickets. Click the ticket you sent in M25.",
        "Canned fix: choose \"Cashier locked out\". Status: pending. Press Add."),
     "Each open ticket shows time left (e.g. \"23h 57m left\"; urgent ones 4h). After Add: the note \"Five wrong PINs lock that person for 5 minutes…\" "
     "appears with \"fix: Cashier locked out · → pending\", the list shows the ticket as pending with SLA \"answered\".", True),
    ("A11", "Canned fix that presses a button on the register", steps(
        "Open the demo ticket \"Receipts come out faint\" (Register 1).",
        "Canned fix \"Printer out of paper (sends Printer test)\", press Add."),
     FLAG_TEXT + "The note shows \"remote action sent\". (Register 1 isn't connected, so the action waits in its queue; that's expected.)", False),
    ("A12", "Hardware swap to repair (RMA)", steps(
        "Open Hardware. On DEMO-PRI-0003 (Register 3) press …",
        "Choose DEMO-PRI-SPARE, reason Faint print, press \"Swap (old one to RMA)\".",
        "On DEMO-PRI-0003 press … → Back from repair."),
     FLAG_TEXT + "After the swap DEMO-PRI-0003 is \"rma\" and DEMO-PRI-SPARE is installed at Register 3. After Back from repair, DEMO-PRI-0003 is \"in stock\". "
     "The small history line under the unit lists each move.", False),
    ("A13", "Fleet: remote action (see also X5)", steps("Fleet → Register 3 (new) → Sign cashier out."),
     FLAG_TEXT + "The action's status becomes succeeded (\"Signed out …\") within a few seconds, and the register goes back to \"Who's working?\".", False),
    ("A14", "Alerts console with runbooks", steps("Open Alerts."),
     FLAG_TEXT + "Open alerts, each with a grey explanation of what it means; some have a fix button (e.g. Force sync) and all store alerts have \"Open a ticket\".", False),
    ("§", "Platform"),
    ("A15", "Staged rollout and kill switch", steps(
        "Open Rollouts. Under \"Price check mode\" type a reason (test kill switch) and press \"Kill switch: off everywhere\".",
        "Look at the register for 10 seconds.",
        "Back in Rollouts, under Price check mode type a reason (restore) and press \"Lift the kill switch: set stage\"."),
     "The card says \"Killed (off everywhere)\", \"On for 0 of 2 stores\", and a history line with your reason. The register's Price check button disappears "
     "within seconds, and comes back after you lift it.", True),
    ("A16", "Partner API key (shown once)", steps("Partners → API keys: Store Journal Square, name Bookkeeper test, keep sales:read. Create key. Press \"I've copied it\"."),
     FLAG_TEXT + "A box shows a key starting adp_… once. After closing it, the list shows only \"adp_xxxxxxxx_…\" (never the whole key) with a Revoke button.", False),
    ("A17", "Webhook", steps("Partners → Webhooks: Store Journal Square, URL https://example.com/adpay, keep sale.completed, Add webhook. Ring a cash sale on the register. After 30 seconds press Deliveries."),
     FLAG_TEXT + "A signing secret (whsec_…) is shown once. Deliveries lists a sale.completed row; since example.com isn't a real partner it will show pending or failed with an HTTP code (that's fine). "
     "Turn the webhook off afterwards (\"Turn off\").", False),
    ("A18", "Agent / referral partner", steps("Agents → Add an agent: name Test Agent, Referral partner, code TEST01, from the 1st of this month, Share of revenue, split 20, bounty 50. Add."),
     FLAG_TEXT + "Test Agent appears with code TEST01 and \"20% of revenue + $50.00 per store going live\". Statements for this month show nothing until a store is assigned (A7).", False),
    ("A19", "Investor / bank pack", steps("Open Investor pack."),
     "\"American Dream Pay — company snapshot\" with Stores live, Registers paired, Sales this month, Revenue, Retention, a 12-month table (By month), Cohorts and Notes. "
     "Revenue shows $0.00 because the demo stores have no pricing plan (see Don't Report This). \"Print / save as PDF\" opens the print dialog.", True),
    ("A20", "UPC library", steps("Open UPC library. Tick \"Names disagree\", then untick it. Search for tissues."),
     "A summary line (\"N product barcodes across all stores…\") and a table including 036000291452 · Test Tissues 2-Ply · 1 store (from Merchant case M31). The search finds it.", True),
    ("A21", "Money: KPIs and statement analyzer", steps(
        "Money → KPIs, then Statement analyzer.",
        "Enter: Store Test Bodega, Processor Clover, Card volume 42000, Transactions 2800, Total card fees 1386, Interchange 1008, POS per month 79, Registers 2, "
        "Dual pricing % 4, Subscription 49.00."),
     FLAG_TEXT + "They pay today $1,465.00, effective rate 3.3%, markup over interchange 0.9% ($378.00); the dual-pricing offer \"Costs $49.00 a month · saves $1,416.00 a month ($16,992.00 a year)\".", False),
    ("A22", "Audit log", steps("Open Audit log."),
     FLAG_TEXT + "Recent entries for your actions (price change, rollout, ticket note, …) with who and when.", False),
    ("§", "New"),
    ("A23", "Give an item a PLU (needed for Register R47)", steps("Merchants → Journal Square → Catalog (Jersey City). Open Banana, type 4011 in PLU, press Save changes."),
     FLAG_TEXT + "Saved; the item list shows Banana with \"· PLU 4011\". Within 15 seconds the register can ring it with the PLU key.", False),
    ("A24", "Move from NRS: the same file from Admin (after Merchant M36)", steps(
        "Merchants → Journal Square → Catalog. Open the \"Move from NRS\" panel.",
        "Choose nrs-sample.json and press Preview. Don't import."),
     FLAG_TEXT + "\"Preview (nothing saved yet): Will import 12 of 12 items: 0 new, 0 updated, 12 unchanged.\" with the same department table and \"Couldn't carry over exactly\" list as the merchant app.", False),
]

CROSS = [
    ("X1", "A price change in Admin reaches an open register", steps(
        "Keep the register open on Snacks (don't reload).",
        "In Admin do A4 (Snickers → 2.09).",
        "Afterwards set Snickers back to 1.99 in Admin."),
     FLAG_TEXT + "Within about 15 seconds the register's Snickers tile shows $2.09 without reloading.", False),
    ("X2", "A sale appears in the owner's live ticker", steps(
        "Keep the merchant app on Sales → Today (LIVE card \"● connected\").",
        "On the register ring Hot Coffee — Medium and tap Exact on the keypad."),
     FLAG_TEXT + "Within about 5 seconds a new top row in LIVE: the time, \"Register 3 (new) · Jersey City\", the cashier, \"cash\", $2.40.", False),
    ("X3", "Closing the drawer reaches the owner", steps("After R25, open the merchant app → Cash → Today."),
     "The same short amount as the register showed in R25, under OVER / SHORT, BY CASHIER (Maria Santos) and DRAWER SESSIONS. The closed drawer is not under \"In the drawers now\".", True),
    ("X4", "An alert fires and arrives in both apps", steps("After R25 (short by more than $5), wait up to 2 minutes."),
     "Merchant app → Alerts shows \"Register 3 (new): drawer counted $Y short\" (Y from R25). The same alert is in Admin → Alerts.", True),
    ("X5", "A remote action from Admin takes effect on the register", steps(
        "Sign in on the register as anyone.",
        "Admin → Fleet → Register 3 (new) → Sign cashier out."),
     FLAG_TEXT + "Within a few seconds the register shows \"Who's working?\" without anyone touching it.", False),
    ("X6", "A kill switch in Admin reaches the register", steps("Same as Admin case A15."),
     "The register's Price check button disappears within seconds and comes back when the switch is lifted.", True),
    ("X7", "An owner's deal reaches the register", steps("Same as Merchant case M18."),
     "Two Beef Sticks ring at $2.00 on the register within about 15 seconds of saving the deal.", True),
    ("X8", "The product library across two stores", steps(
        "Merchant app as Kevin (2015550142): Items → + Add item, name Test Water 1L, category Drinks, cash 1.25, barcode 012345678905. Save.",
        "On the register (Journal Square) type 012345678905 in the search box and press Enter."),
     FLAG_TEXT + "The register's \"New item\" box opens with \"Test Water 1L\" already in the name box and the line \"Other stores call it “Test Water 1L”.\" "
     "The price box stays empty (each store sets its own). Press Cancel.", False),
    ("X9", "A vendor order is received at the register", steps(
        "Do M21 first (order sent to Tuscan Dairy).",
        "On the register press Receive. Under \"Against an order?\" tap Tuscan Dairy. Add Whole Milk — Gallon (type whole milk), set the number, press \"Add … to stock\".",
        "In the merchant app open Orders."),
     "Receive shows \"Against an order? Tuscan Dairy\". " + FLAG_TEXT.strip() + " After receiving, the order in the merchant app shows \"received N\" next to \"ordered N\".", False),
    ("X10", "A checklist on the register reaches the owner", steps("Do R40, then merchant app → Hours → Open & close checklists."),
     "Today's Jersey City row shows Closing \"2/5 · 2 photos missing\" with the time and Maria Santos.", True),
    ("X11", "A loyalty sign-up reaches the customer list", steps("Do R31–R32, then merchant app → Customers."),
     "The CUSTOMERS list includes \"Phone ···0123\" with a visit and the amount spent.", True),
    ("X12", "A support ticket goes both ways", steps(
        "Do M25 (store sends a problem), then A10 (AD Pay answers with a canned fix).",
        "In the merchant app, Help → Problems & equipment, tap your ticket."),
     FLAG_TEXT + "The ticket says \"Waiting for you\"; tapping it shows AD Pay's reply (the locked-out explanation). Reply \"Thanks\": in Admin the ticket goes back to open.", False),
    ("X13", "An owner's key page reaches the register", steps("Do Merchant MP1, then Register R48."),
     FLAG_TEXT + "The Coffee bar tab appears on the register within about 15 seconds of saving, without reloading.", False),
    ("X14", "An NRS import scans on the register straight away", steps(
        "After Merchant M36, wait 15 seconds. On the register look at the tabs.",
        "Click the search box, type 049000028911 and press Enter (that's \"scanning\" the Coca-Cola barcode).",
        "Type 028200003843 and press Enter. Press \"ID checked — 21+\".",
        "Void the ticket."),
     FLAG_TEXT + "New tabs: Grocery tax, Grocery Non-Taxable, Beers 21+, all smoke 21+. The Coca-Cola barcode rings Coca-Cola Classic 20oz $2.49 (Tax $0.16). "
     "The Marlboro barcode asks for the 21+ check, then rings Marlboro Red Box $13.99 with its tax inside: the ticket shows Subtotal $16.48, Tax $0.16, "
     "\"Tax included in prices $0.87\", Cash $16.64.", False),
    ("X15", "\"Price includes tax\" reaches the register", steps("Do Merchant M33, then Register R50."),
     FLAG_TEXT + "Within 15 seconds of saving, Hot Coffee — Large rings with no tax added on top and the tax shown as \"included\".", False),
]

DONT = [
    ("Stubbed on purpose (no outside account yet)", [
        "Card payments are simulated. They're approved, except any card amount ending in .13, which always declines on purpose.",
        "No text message or email is ever sent: sign-in codes are always 123456 and shown on screen; \"text me my receipt\", promo texts and vendor orders say they were recorded but not delivered.",
        "No push notifications; alerts show only inside the apps.",
        "Processor costs, deposits and fees are typed by hand or show \"once the processor account is live\". KYB shows \"not started\".",
        "Revenue and margin show $0.00 or — because the demo stores have no pricing plan and no processor cost entered.",
        "Webhooks only reach a real partner server; example.com will show pending/failed deliveries.",
        "Paying agents happens outside the system; Agents only shows the statement.",
        "Check / other: EBT, gift cards and house accounts only record the payment and its reference. No EBT processor, gift-card balance or account balance is checked.",
        "A refund without a receipt is paid back in cash only.",
        "Getting a store's file out of NRS needs that store's own NRS login, so you use the made-up sample file instead.",
    ]),
    ("Needs hardware we don't have here", [
        "No real printer, cash drawer or barcode scanner: receipts appear as \"Receipt (printer preview)\", the drawer as a \"Drawer opened\" message, and \"scanning\" is typing the barcode + Enter.",
        "Photos (checklist, count sheet, item photos, documents) use your computer's file picker instead of a camera.",
        "Driver's-licence ID scanning can't be tested.",
        "Printing labels straight to a label printer isn't possible; tags come out as a PDF.",
        "Admin → Fleet: \"Re-pair terminal\", \"Roll back build\" and \"Reboot device\" are greyed out on purpose.",
    ]),
    ("Waiting on a decision or a contract", [
        "Rolling out new versions of the register app to stores (over-the-air updates) isn't built; Rollouts only switches features on and off.",
        "Only English and Spanish are offered to customers; the other languages are drafts until a translator reviews them.",
        "The \"text me deals\" consent wording is shown in English in every language (legal wording, pending review).",
        "Tax & compliance templates say \"draft\": starting values, not legal advice.",
        "The product (UPC) library only knows what our stores have typed in; a full product database needs a licence.",
    ]),
    ("Dropped", [
        "Lottery features were dropped at the founder's request. The Lottery category and its items are ordinary items; there's no lottery-specific screen or report.",
    ]),
    ("Known small issues (already on our list)", [
        "The deal line on the customer screen (\"2 for $6.00 · Energy drinks\") stays in English when the customer picks Spanish.",
    ]),
    ("Not bugs", [
        "The first time a page opens it can take 10–30 seconds, and everything is a little slow: it runs on one laptop over the internet.",
        "The data has months of history and other people may be testing at the same time: totals don't start at zero, and you may see sales and tickets that aren't yours.",
        "Things you add while testing (Test Muffin, Test Salsa, the Coffee bar page, the NRS sample items and their tabs) stay in the demo afterwards; other testers may see them too.",
        "Numbers that depend on time or on other people (stock counts, \"expires in N days\", visit counts, SLA hours left) are described as \"varies\" in the cases.",
    ]),
]


def header_row(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER


def row_height(cells: list[tuple[str, int]]) -> float:
    """Excel doesn't auto-size wrapped rows in a generated file: estimate lines per cell from the column width, generously."""

    def lines(text: str, width: int) -> int:
        per_line = max(10, int(width * 0.95))
        return sum(max(1, -(-len(p) // per_line)) for p in text.split("\n"))

    most = max(lines(t, w) for t, w in cells)
    return min(max(28, most * 13.5 + 8), 409)


def case_sheet(wb: Workbook, name: str, intro: str, cases: list) -> None:
    ws = wb.create_sheet(name)
    ws["A1"] = name
    ws["A1"].font = TITLE
    ws["A2"] = intro
    ws["A2"].font = MUTED
    ws.merge_cells("A2:F2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 42
    header_row(ws, 3, ["ID", "Feature", "Steps", "Expected Result", "Pass/Fail", "Notes"])
    widths = [7, 26, 58, 64, 12, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    dv = DataValidation(type="list", formula1='"Pass,Fail,Skipped"', allow_blank=True, showDropDown=False)
    dv.error = "Choose Pass, Fail or Skipped"
    dv.prompt = "Pass, Fail or Skipped"
    ws.add_data_validation(dv)
    row = 4
    for case in cases:
        if case[0] in ("§", "§!"):
            c = ws.cell(row=row, column=1, value=case[1])
            c.font = BOLD
            c.alignment = WRAP
            for col in range(1, 7):
                ws.cell(row=row, column=col).fill = SCREEN_FILL if case[0] == "§!" else SECTION_FILL
                ws.cell(row=row, column=col).border = BORDER
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            if case[0] == "§!":
                ws.row_dimensions[row].height = 34
            row += 1
            continue
        cid, feature, step_text, expected, checked = case
        if checked == "screen":
            expected = SCREEN_TEXT + expected
        values = [cid, feature, step_text, expected, None, None]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = BOLD if col == 1 else BODY
            c.alignment = WRAP
            c.border = BORDER
        if checked == "screen":
            for col in (1, 2, 4):
                ws.cell(row=row, column=col).fill = SCREEN_FILL
        elif not checked:
            ws.cell(row=row, column=1).fill = FLAG_FILL
            ws.cell(row=row, column=4).fill = FLAG_FILL
        ws.cell(row=row, column=5).fill = FILL_IN
        ws.cell(row=row, column=6).fill = FILL_IN
        dv.add(ws.cell(row=row, column=5))
        ws.row_dimensions[row].height = row_height([(feature, 26), (step_text, 58), (expected, 64)])
        row += 1
    ws.freeze_panes = "A4"
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


wb = Workbook()
start = wb.active
start.title = "Start Here"
start.column_dimensions["A"].width = 30
start.column_dimensions["B"].width = 62
start.column_dimensions["C"].width = 44
r = 1
start.cell(row=r, column=1, value="AD Pay POS — test plan").font = TITLE
r += 2
about = [
    "AD Pay is a point-of-sale system for delis and corner stores. There are three apps: the Register (the till at the counter), the Merchant app "
    "(the store owner's phone app, used here in a web browser) and Admin (AD Pay's own back office).",
    "Everything is demo data. Nothing costs money and no real card is ever charged. Use Chrome on a computer, one tab per app.",
    "Work through the sheets Register → Merchant App → Admin → Cross-App, top to bottom (some cases build on earlier ones). In each row, follow the Steps, "
    "compare with Expected Result, and pick Pass, Fail or Skipped in the yellow Pass/Fail cell. A case passes only if what you see matches.",
    "At the top of the Register and Merchant App sheets there are LIGHT RED rows: whole screens we have never clicked through (only automated tests). Please do those first "
    "(after signing in) and look at them carefully. Rows with an amber ID are ones we worked out from the code without clicking through ourselves: if they don't match, report them anyway. "
    "For every Fail, add a row on the Bug Reports sheet. Before reporting, check the \"Don't Report These\" sheet.",
]
for line in about:
    c = start.cell(row=r, column=1, value=line)
    c.font = BODY
    c.alignment = WRAP
    start.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    start.row_dimensions[r].height = 44
    r += 1
r += 1


def section(title: str) -> None:
    global r
    start.cell(row=r, column=1, value=title).font = H2
    r += 1


def pair_rows(rows: list[tuple], headers: list[str] | None = None) -> None:
    global r
    if headers:
        header_row(start, r, headers)
        r += 1
    for values in rows:
        for col, v in enumerate(values, 1):
            c = start.cell(row=r, column=col, value=v)
            c.font = BODY
            c.alignment = WRAP
            c.border = BORDER
        r += 1
    r += 1


section("1. The three links (click to open). They change if the system is restarted: if one doesn't open, ask for new links.")
for label, url in [("Register", REGISTER_URL), ("Merchant app", MERCHANT_URL), ("Admin", ADMIN_URL)]:
    start.cell(row=r, column=1, value=label).font = BOLD
    c = start.cell(row=r, column=2, value=url)
    c.hyperlink = url
    c.font = LINK
    r += 1
r += 1
section("2. Admin sign-in")
pair_rows([("Email", "admin@adpay.local"), ("Password", "adpay-demo")])
section("3. Register setup codes (a code connects the register to a store; use JSQ3-DEMO unless a case says otherwise)")
pair_rows(
    [
        ("JSQ3-DEMO", "Journal Square Deli & Grocery · Jersey City · Register 3 (new)", "Use this one"),
        ("JSQ1-DEMO", "Journal Square Deli & Grocery · Jersey City · Register 1", ""),
        ("JSQ2-DEMO", "Journal Square Deli & Grocery · Jersey City · Register 2", ""),
        ("AST1-DEMO", "Journal Square Deli & Grocery · Astoria · Register 1", ""),
        ("BAY1-DEMO", "Bayonne Corner Mart · Broadway · Register 1", "A separate business"),
    ],
    ["Code", "Store it pairs", ""],
)
section("4. Merchant app sign-in (no text message is sent: the code is always 123456, and it's also shown on screen)")
pair_rows(
    [
        ("2015550100", "Nadia Haddad, owner", "Journal Square Deli & Grocery"),
        ("2015550101", "Luis Ortega, manager", "Journal Square Deli & Grocery"),
        ("2015550142", "Kevin Walsh, owner", "Bayonne Corner Mart (a separate business)"),
    ],
    ["Phone", "Person", "Store"],
)
section("5. Register PINs (at \"Who's working?\" tap the name, then tap the 4 digits: it signs in on the last digit, there is no Enter)")
pair_rows(
    [
        ("2580", "Nadia Haddad · Owner", "Journal Square"),
        ("1357", "Luis Ortega · Manager", "Journal Square"),
        ("2468", "Maria Santos · Cashier", "Journal Square"),
        ("3690", "Dev Patel · Cashier", "Journal Square"),
        ("2580", "Kevin Walsh · Owner", "Bayonne Corner Mart"),
        ("4826", "Aisha Khan · Cashier", "Bayonne Corner Mart"),
    ],
    ["PIN", "Name · Role", "Store"],
)
section("6. Money in the Jersey City store")
for line in [
    "Most items are taxed at 6.625%. Grocery, Tobacco and Lottery items are not taxed.",
    "The register's keypad is always on screen: type an amount, then press a tender (Cash, Card, Check / other) or a key that uses the number (@, PLU, → Department).",
    "The card price is 4% higher than the cash price (tobacco and lottery: same price). All the amounts in the cases follow from this.",
    "Owners and managers can approve things a cashier isn't allowed to do: the register asks \"Who's approving?\".",
]:
    c = start.cell(row=r, column=1, value=line)
    c.font = BODY
    c.alignment = WRAP
    start.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    start.row_dimensions[r].height = 18
    r += 1
start.sheet_view.showGridLines = False

case_sheet(wb, "Register", "Do these in order on one register, Journal Square · Jersey City · Register 3. The register is built for a wide screen: make the Chrome window large.", REGISTER)
case_sheet(wb, "Merchant App", "Sign in as the owner Nadia (2015550100) unless a case says otherwise. Keep the register open in another tab: several cases check it.", MERCHANT)
case_sheet(wb, "Admin", "Sign in with admin@adpay.local / adpay-demo. Some cases feed register or merchant cases (they say which).", ADMIN)
case_sheet(wb, "Cross-App", "Flows that go from one app to another. Keep the apps side by side. Most need a case from another sheet done first.", CROSS)

ws = wb.create_sheet("Don't Report These")
ws["A1"] = "Don't report these (known and expected)"
ws["A1"].font = TITLE
ws.column_dimensions["A"].width = 36
ws.column_dimensions["B"].width = 110
row = 3
for title, items in DONT:
    c = ws.cell(row=row, column=1, value=title)
    c.font = BOLD
    c.fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1
    for item in items:
        ws.cell(row=row, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
        c = ws.cell(row=row, column=2, value=item)
        c.font = BODY
        c.alignment = WRAP
        ws.row_dimensions[row].height = max(16, (len(item) // 105 + 1) * 15)
        row += 1
    row += 1
ws.sheet_view.showGridLines = False

nrs = wb.create_sheet("NRS sample file")
nrs["A1"] = "NRS sample file (for Merchant cases M34–M37)"
nrs["A1"].font = TITLE
nrs.column_dimensions["A"].width = 150
for i, line in enumerate([
    "A made-up 12-item price book in the exact shape a store's NRS portal gives us. To save it as a file:",
    "1. Click cell A9 (the long line of text below). 2. Press Ctrl+C (Mac: Cmd+C).",
    "3. Windows: open Notepad, paste (Ctrl+V), File → Save as, set \"Save as type\" to All files, name it nrs-sample.json. "
    "Mac: open TextEdit, Format → Make Plain Text, paste (Cmd+V), save as nrs-sample.json (if it asks, use .json).",
    "4. Check the file starts with [{ and ends with }]. Don't change anything inside it.",
    "The same file is in the project as docs/nrs-sample-pricebook.json if someone can send it to you.",
], 3):
    c = nrs.cell(row=i, column=1, value=line)
    c.font = BODY
    c.alignment = WRAP
import json as _json
_sample = _json.loads((ROOT / "docs" / "nrs-sample-pricebook.json").read_text(encoding="utf-8"))
c = nrs.cell(row=9, column=1, value=_json.dumps(_sample, ensure_ascii=False, separators=(",", ":")))
c.font = Font(name="Consolas", size=9)
c.fill = FILL_IN
nrs.sheet_view.showGridLines = False

bugs = wb.create_sheet("Bug Reports")
bugs["A1"] = "Bug reports"
bugs["A1"].font = TITLE
bugs["A2"] = ("One row per problem. Fill in the yellow cells. Example: Test ID R17 · App Register · What I did: charged $9.13 by card · "
              "What I expected: the decline message · What happened: the sale completed · Time: Sep 27, 3:42 pm ET · Screenshot: yes, sent by email.")
bugs["A2"].font = MUTED
bugs["A2"].alignment = WRAP
bugs.merge_cells("A2:G2")
bugs.row_dimensions[2].height = 32
header_row(bugs, 3, ["Test ID", "App", "What I did", "What I expected", "What happened", "Time", "Screenshot"])
for i, w in enumerate([10, 16, 44, 40, 44, 22, 18], 1):
    bugs.column_dimensions[chr(64 + i)].width = w
apps = DataValidation(type="list", formula1='"Register,Merchant App,Admin,Cross-App"', allow_blank=True)
bugs.add_data_validation(apps)
for rr in range(4, 54):
    for col in range(1, 8):
        c = bugs.cell(row=rr, column=col)
        c.fill = FILL_IN
        c.border = BORDER
        c.font = BODY
        c.alignment = WRAP
    apps.add(bugs.cell(row=rr, column=2))
bugs.freeze_panes = "A4"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
cases = [c for sh in (REGISTER, MERCHANT, ADMIN, CROSS) for c in sh if c[0] not in ("§", "§!")]
print("cases:", len(cases), "worked out from code:", sum(1 for c in cases if c[4] is False), "never-clicked screens:", sum(1 for c in cases if c[4] == "screen"))
